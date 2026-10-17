"""Independent full-source recomputation and read-only inspection of delivered Excel."""

import csv
import json
import sys
import zipfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from sales import write_json

a = json.loads((ROOT / "evidence/analysis.json").read_text())
m = a["metrics"]
with (ROOT / "data/raw/train.csv").open(encoding="utf-8-sig", newline="") as stream:
    source = list(csv.DictReader(stream))
monthly = {}
incoming = {}
for row in source:
    month = datetime.strptime(row["Order Date"], "%d/%m/%Y").strftime("%Y-%m")
    monthly[month] = monthly.get(month, 0) + int(Decimal(row["Sales"]) * 10000)
for path in sorted((ROOT / "data/incoming").glob("*.csv")):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        for row in csv.DictReader(stream):
            incoming[row["Row ID"]] = row
checks = {
    "source_rows": len(source) == m["lines"] == 9800,
    "source_orders": len({row["Order ID"] for row in source}) == m["orders"],
    "exact_total": sum(int(Decimal(row["Sales"]) * 10000) for row in source)
    == m["sales_units"],
    "monthly_totals": monthly
    == {x["month"]: x["sales_units"] for x in a["reports"]["monthly"]},
    "all_rows_retained": len(incoming) == len(source),
    "selected_source_values_retained": all(
        all(
            incoming[row["Row ID"]][field] == row[field]
            for field in incoming[row["Row ID"]]
        )
        for row in source
    ),
    "no_customer_names_in_deliveries": all(
        "Customer Name" not in row for row in incoming.values()
    ),
    "all_six_reconciliations_pass": all(a["checks"].values()),
}
book = ROOT / "deliverables/monthly-sales_2018-12.xlsx"
values = load_workbook(book, data_only=True)
formulas = load_workbook(book, data_only=False)
s = values["Monthly report"]
f = formulas["Monthly report"]
end = len(a["reports"]["monthly"]) + 5
checks.update(
    {
        "cached_total_matches": abs(s["B11"].value - m["sales_units"] / 10000) < 1e-7,
        "cached_last_month_matches": abs(s["B7"].value - monthly["2018-12"] / 10000)
        < 1e-7,
        "cached_yoy_matches": abs(
            s["B10"].value - (monthly["2018-12"] / monthly["2017-12"] - 1)
        )
        < 1e-10,
        "report_month_is_typed_date": isinstance(s["B6"].value, datetime)
        and s["B6"].value.strftime("%Y-%m") == "2018-12",
        "summary_linked_to_calculations": f["B7"].data_type == "f"
        and "Monthly detail" in f["B7"].value,
        "first_yoy_unavailable": values["Monthly detail"]["I6"].value == "n.a.",
        "last_yoy_formula_exists": formulas["Monthly detail"][f"I{end}"].data_type
        == "f",
        "native_chart_exists": len(formulas["Monthly report"]._charts) == 1,
        "review_status_dropdown": len(
            formulas["Category review"].data_validations.dataValidation
        )
        == 1,
        "review_note_keys_are_text": values["Category review"]["A6"].value == "2015-01",
        "delivered_notes_are_blank": all(
            not values["Category review"].cell(row, c).value
            for row in range(6, 150)
            for c in [4, 6]
        ),
    }
)
with zipfile.ZipFile(book) as z:
    charts = [
        z.read(n).decode()
        for n in z.namelist()
        if "/charts/" in n and n.endswith(".xml")
    ]
    chart = next(xml for xml in charts if "lineChart" in xml)
    checks["three_native_charts"] = len(charts) == 3
    checks["comparison_chart_palette"] = (
        sum("126F69" in xml.upper() and "B8C1C7" in xml.upper() for xml in charts) == 2
    )
    checks["chart_bound_to_cells"] = f"$M${end+4}:$M${end+15}" in chart
    checks["chart_line_color_exported"] = "126F69" in chart.upper()
demo = json.loads((ROOT / "docs/refresh-demonstration.json").read_text())
checks["real_workbook_refresh_demonstration"] = demo["all_passed"]
from pypdf import PdfReader

# Recompute matched business findings directly from original rows, independently of SQL.
br = a["business_review"]
bs = br["summary"]
current_rows = [
    r
    for r in source
    if datetime.strptime(r["Order Date"], "%d/%m/%Y").strftime("%Y-%m-%d")
    >= bs["current_start"]
    and datetime.strptime(r["Order Date"], "%d/%m/%Y").strftime("%Y-%m-%d")
    <= bs["current_end"]
]
prior_rows = [
    r
    for r in source
    if datetime.strptime(r["Order Date"], "%d/%m/%Y").strftime("%Y-%m-%d")
    >= bs["prior_start"]
    and datetime.strptime(r["Order Date"], "%d/%m/%Y").strftime("%Y-%m-%d")
    <= bs["prior_end"]
]
amount = lambda rows: sum(int(Decimal(r["Sales"]) * 10000) for r in rows)
checks["matched_source_sales"] = (
    amount(current_rows) == bs["current"]["sales_units"]
    and amount(prior_rows) == bs["prior"]["sales_units"]
)
checks["matched_source_orders"] = (
    len({r["Order ID"] for r in current_rows}) == bs["current"]["orders"]
    and len({r["Order ID"] for r in prior_rows}) == bs["prior"]["orders"]
)
checks["matching_day_cutoffs"] = (
    bs["current_end"] == "2018-12-30" and bs["prior_end"] == "2017-12-30"
)
expected_actions = []
for scope, field in [("category", "Category"), ("region", "Region")]:
    for driver in br["drivers"][scope]:
        cur = [r for r in current_rows if r[field] == driver["member"]]
        prev = [r for r in prior_rows if r[field] == driver["member"]]
        checks[scope + "_" + driver["member"] + "_source_delta"] = (
            amount(cur) - amount(prev) == driver["change_units"]
        )
        if amount(cur) < amount(prev):
            products = {r["Product ID"] for r in cur + prev}
            losses = [
                (
                    amount([r for r in cur if r["Product ID"] == product])
                    - amount([r for r in prev if r["Product ID"] == product]),
                    product,
                )
                for product in products
            ]
            ids = ", ".join(
                product for delta, product in sorted(losses) if delta < 0
            ).split(", ")[:3]
            expected_actions.append(
                (amount(cur) - amount(prev), scope, driver["member"], ", ".join(ids))
            )
checks["review_rank_and_source_product_ids"] = [
    (r["sales_change"], r["scope"], r["segment"], r["product_ids_to_check"])
    for r in br["actions"]
] == [
    (delta / 10000, scope, member, ids)
    for delta, scope, member, ids in sorted(expected_actions)
]
checks["all_business_reconciliations"] = len(br["checks"]) == 7 and all(
    br["checks"].values()
)
checks["business_sheet_first"] = values.sheetnames[0] == "Sales review"
checks["business_cached_yoy"] = (
    abs(
        values["Sales review"]["B11"].value
        - (amount(current_rows) / amount(prior_rows) - 1)
    )
    < 1e-10
)
checks["dashboard_card_linked"] = (
    formulas["Sales review"]["B9"].value == "='Monthly detail'!P6"
)
checks["dashboard_card_matches_source"] = (
    abs(values["Sales review"]["B9"].value - amount(current_rows) / 10000) < 1e-7
)
checks["business_bridge_reconciles"] = (
    abs(
        values["Monthly detail"]["P12"].value
        + values["Monthly detail"]["P13"].value
        - (amount(current_rows) - amount(prior_rows)) / 10000
    )
    < 1e-7
)
checks["useful_deliverables_exist"] = all(
    (ROOT / "deliverables" / name).is_file()
    for name in ["sales-review.html", "review_actions.csv", "excel-sales-review.png"]
)
checks["html_reports_actual_findings"] = br["headline"] in (
    ROOT / "deliverables/sales-review.html"
).read_text(encoding="utf-8")

pdf = PdfReader(ROOT / "deliverables/case-study.pdf", strict=True)
checks["pdf_two_pages"] = len(pdf.pages) == 2
checks["pdf_source_identifier"] = (
    "rohitsahoo/sales-forecasting" in pdf.pages[1].extract_text()
)
result = {
    "method": "Independent original-source Decimal/date recomputation; selected-column equality; read-only openpyxl/XML; real-engine refresh replay",
    "checks": checks,
    "all_passed": all(checks.values()),
    "nativeExcelRecalculationTested": False,
}
write_json(ROOT / "docs/independent-verification.json", result)
print(json.dumps(result, indent=2))
if not result["all_passed"]:
    raise SystemExit(1)
